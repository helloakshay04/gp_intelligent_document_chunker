from pathlib import Path
from typing import Any
from html.parser import HTMLParser
import csv, io, re, json, xml.etree.ElementTree as ET
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS={".pdf",".txt",".md",".csv",".xls",".xlsx",".docx",".json",".xml",".html",".py",".java",".js",".sql"}
MAPPING={".pdf":("pdf","application/pdf"),".txt":("text","text/plain"),".md":("markdown","text/markdown"),".csv":("csv","text/csv"),".xls":("spreadsheet","application/vnd.ms-excel"),".xlsx":("spreadsheet","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),".docx":("word","application/vnd.openxmlformats-officedocument.wordprocessingml.document"),".json":("json","application/json"),".xml":("xml","application/xml"),".html":("html","text/html"),".py":("code","text/x-python"),".java":("code","text/x-java-source"),".js":("code","text/javascript"),".sql":("code","application/sql")}
class _HTMLText(HTMLParser):
    def __init__(self): super().__init__(convert_charrefs=True); self.parts=[]; self.skip=0; self.table_depth=0; self.row=[]; self.cell=""; self.rows=[]
    def handle_starttag(self,tag,attrs):
        tag=tag.lower()
        if tag in {"script","style","noscript"}: self.skip+=1
        if self.skip: return
        if tag=="table": self.table_depth+=1
        if tag in {"td","th"}: self.cell=""
        if tag=="tr": self.row=[]
        if tag in {"p","div","section","article","h1","h2","h3","h4","h5","h6","li","br"}: self.parts.append("\n")
    def handle_endtag(self,tag):
        tag=tag.lower()
        if tag in {"script","style","noscript"} and self.skip: self.skip-=1; return
        if self.skip:return
        if tag in {"td","th"}:
            if self.cell.strip(): self.row.append(re.sub(r"\s+"," ",self.cell).strip())
            self.cell=""
        if tag=="tr" and self.row:
            row_text=" | ".join(self.row); self.rows.append((self.table_depth,row_text)); self.parts.append("\n| "+row_text+" |\n"); self.row=[]
        if tag=="table": self.table_depth=max(0,self.table_depth-1); self.parts.append("\n")
        if tag in {"p","div","section","article","h1","h2","h3","h4","h5","h6","li"}: self.parts.append("\n")
    def handle_data(self,data):
        if not self.skip:
            if self.table_depth:
                self.cell+=data
            else:self.parts.append(data)

def detect_file_type(filename: str, content: bytes|None=None)->dict[str,str]:
    ext=Path(filename).suffix.lower()
    if ext not in MAPPING: raise ValueError(f"Unsupported file type '{ext or '(no extension)'}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    if content is not None:
        if ext==".pdf" and not content.startswith(b"%PDF"): raise ValueError("File extension is .pdf but file signature is not a PDF.")
        if ext==".xlsx" and content[:2]!=b"PK": raise ValueError("File extension is .xlsx but file signature is not a ZIP-based Office document.")
        if ext==".docx" and content[:2]!=b"PK": raise ValueError("File extension is .docx but file signature is not a ZIP-based Office document.")
    return {"extension":ext,"file_type":MAPPING[ext][0],"content_type":MAPPING[ext][1]}
def _decode(content:bytes)->str:
    for enc in ("utf-8-sig","utf-8","cp1252","latin-1"):
        try:return content.decode(enc)
        except UnicodeDecodeError:pass
    return content.decode("utf-8",errors="replace")
def _assemble(sections):
    clean=[]; pos=0
    for s in sections:
        txt=s.get("text","").strip()
        if not txt: continue
        if clean: pos+=2
        start=pos; end=start+len(txt); meta=dict(s.get("metadata",{})); meta["start_char"]=start; meta["end_char"]=end
        clean.append({"text":txt,"metadata":meta}); pos=end
    return "\n\n".join(s["text"] for s in clean),clean
def _walk_docx_table(table, sections, table_no):
    for ri,row in enumerate(table.rows,1):
        vals=[]
        for cell in row.cells:
            txt=cell.text.strip()
            nested=list(getattr(cell,"tables",[]))
            if txt: vals.append(txt)
            for nti,nt in enumerate(nested,1): _walk_docx_table(nt,sections,f"{table_no}.{nti}")
        if any(vals): sections.append({"text":" | ".join(vals),"metadata":{"table":table_no,"row":ri,"source_type":"docx_table_row"}})
def extract_document(filename:str,content:bytes)->dict[str,Any]:
    detected=detect_file_type(filename,content); ext=detected["extension"]; warnings=[]; sections=[]; page_count=None; sheet_names=[]
    try:
        if ext==".pdf":
            reader=PdfReader(io.BytesIO(content),strict=False); page_count=len(reader.pages)
            for i,page in enumerate(reader.pages,1):
                try: txt=(page.extract_text() or "").strip()
                except Exception as e: warnings.append(f"Page {i} extraction failed ({type(e).__name__})."); txt=""
                if not txt: warnings.append(f"Page {i} has no selectable text; scanned PDFs require OCR.")
                sections.append({"text":txt,"metadata":{"page_number":i,"source_type":"pdf_page"}})
        elif ext==".csv":
            rows=list(csv.reader(io.StringIO(_decode(content))))
            if rows:
                header=rows[0]
                sections.append({"text":" | ".join(str(x) for x in header),"metadata":{"row_number":1,"source_type":"csv_header"}})
                for idx,row in enumerate(rows[1:],2):
                    cells=[f"{header[j] if j<len(header) and header[j] else f'column_{j+1}'}: {v}" for j,v in enumerate(row) if v]
                    if cells: sections.append({"text":" | ".join(cells),"metadata":{"row_number":idx,"source_type":"csv_row"}})
        elif ext==".xlsx":
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True); sheet_names=list(wb.sheetnames)
            for ws in wb.worksheets:
                rows=ws.iter_rows(values_only=True); header=next(rows,None)
                if header is None: continue
                header=[str(x) if x is not None else f"column_{i+1}" for i,x in enumerate(header)]
                sections.append({"text":f"Sheet: {ws.title}\nColumns: {' | '.join(header)}","metadata":{"sheet_name":ws.title,"row_number":1,"source_type":"spreadsheet_header"}})
                for ri,row in enumerate(rows,2):
                    cells=[f"{header[i]}: {v}" for i,v in enumerate(row) if v is not None and i<len(header)]
                    if cells: sections.append({"text":" | ".join(cells),"metadata":{"sheet_name":ws.title,"row_number":ri,"source_type":"spreadsheet_row"}})
            wb.close()
            if not sections:warnings.append("Workbook has no populated cells.")
        elif ext==".xls": raise ValueError("Legacy .xls is recognized but not supported by the built-in extractor. Save it as .xlsx or .csv.")
        elif ext==".docx":
            doc=DocxDocument(io.BytesIO(content))
            for i,p in enumerate(doc.paragraphs,1):
                if p.text.strip(): sections.append({"text":p.text.strip(),"metadata":{"paragraph_number":i,"section":p.style.name if p.style else None,"source_type":"docx_paragraph"}})
            for ti,t in enumerate(doc.tables,1): _walk_docx_table(t,sections,ti)
        else:
            txt=_decode(content)
            if ext==".html":
                parser=_HTMLText(); parser.feed(txt); parser.close()
                body=re.sub(r"[ \t]+"," ","".join(parser.parts)); body=re.sub(r"\n\s*\n\s*\n+","\n\n",body).strip()
                # Flattened body retains tables as text; attach detected row summaries as warnings/signals rather than duplicating rows.
                sections=[{"text":body,"metadata":{"source_type":"html_text","html_table_rows":len(parser.rows),"nested_table_detected":any(depth>1 for depth,_ in parser.rows)}}]
            elif ext==".json":
                try:
                    parsed=json.loads(txt); txt=json.dumps(parsed,ensure_ascii=False,indent=2)
                except json.JSONDecodeError as e: warnings.append(f"JSON parse warning at line {e.lineno}, column {e.colno}; preserved raw text.")
                sections=[{"text":txt,"metadata":{"source_type":"json_text"}}]
            elif ext==".xml":
                try:
                    root=ET.fromstring(txt)
                    def xml_lines(el,depth=0):
                        attrs=" ".join(f"{k}={v}" for k,v in el.attrib.items()); value=(el.text or "").strip(); label=f"{'  '*depth}<{el.tag}{(' '+attrs) if attrs else ''}>{(': '+value) if value else ''}"
                        out=[label]
                        for child in list(el):out.extend(xml_lines(child,depth+1))
                        return out
                    txt="\n".join(xml_lines(root))
                except ET.ParseError as e:warnings.append(f"XML parse warning at line {e.position[0]}; preserved raw text.")
                sections=[{"text":txt,"metadata":{"source_type":"xml_text"}}]
            else: sections=[{"text":txt.strip(),"metadata":{"source_type":"text_file"}}]
            if ext==".md":
                sections=[]
                for i,line in enumerate(txt.splitlines(),1):
                    if line.strip():
                        meta={"line_number":i,"source_type":"markdown_line"}
                        if re.match(r"^\s{0,3}#{1,6}\s",line):meta["section"]=line.lstrip("# ").strip()
                        if re.match(r"^\s*\|.*\|\s*$",line):meta["table_row"]=True
                        sections.append({"text":line.strip(),"metadata":meta})
        text,sections=_assemble(sections)
    except ValueError: raise
    except Exception as e: raise ValueError(f"Unable to extract {ext} document ({type(e).__name__}): {e}") from e
    if len(content)>0 and not text.strip(): warnings.append("No extractable text found; this may be an image-only or empty document.")
    structure_signals={}
    if ext==".html": structure_signals={"html_table_rows":len(parser.rows),"nested_table_detected":any(depth>1 for depth,_ in parser.rows)}
    if ext==".docx": structure_signals={"docx_table_count":len(DocxDocument(io.BytesIO(content)).tables)}
    return {**detected,"text":text,"sections":sections,"page_count":page_count,"sheet_names":sheet_names,"warnings":list(dict.fromkeys(warnings)),"structure_signals":structure_signals}
