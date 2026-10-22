from pathlib import Path
import time, uuid
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Request
from fastapi.responses import FileResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ValidationError
from chunker_sdk import DocumentChunker, ChunkingConfig, __version__
from chunker_sdk.extractors import SUPPORTED_EXTENSIONS
from chunker_sdk.telemetry import span, get_metrics, prometheus_metrics, logger, record_call
from fastapi.middleware.cors import CORSMiddleware

BASE_DIR=Path(__file__).resolve().parents[1]; UI_DIR=BASE_DIR/"ui"; engine=DocumentChunker()
app = FastAPI(
    title="Intelligent Document Chunker",
    docs_url="/docs",
    redoc_url="/redoc",
    openapi_url="/openapi.json",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://subtle-chimera-be1593.netlify.app",
        "http://localhost:4200",
    ],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/ui",StaticFiles(directory=UI_DIR),name="ui")
@app.middleware("http")
async def request_observability(request:Request,call_next):
    trace_id=request.headers.get("x-trace-id") or uuid.uuid4().hex[:16]; request.state.trace_id=trace_id; started=time.perf_counter()
    try:
        with span("http_request",trace_id,method=request.method,path=request.url.path): response=await call_next(request)
        response.headers["X-Trace-Id"]=trace_id; response.headers["X-Process-Time-Ms"]=str(round((time.perf_counter()-started)*1000,3)); return response
    except Exception as exc:
        logger.error("http.request.error",extra={"trace_id":trace_id,"span":"http_request","error_type":type(exc).__name__}); raise
@app.get("/health")
def health():return {"status":"ok","sdk_version":__version__}
@app.get("/metrics")
def metrics():return get_metrics()
@app.get("/metrics/prometheus",response_class=PlainTextResponse)
def metrics_prometheus():return prometheus_metrics()
@app.get("/v1/config")
def config():return {"sdk_version":__version__,"supported_extensions":sorted(SUPPORTED_EXTENSIONS),"bounds":{"chunk_size":{"min":100,"max":10000,"default":1000},"overlap":{"min":0,"max":2000,"default":150,"max_percent_of_chunk_size":50},"preview_count":{"min":5,"max":10,"default":8}},"strategies":["auto","fixed","recursive","semantic","sentence","paragraph","table","code","parent_child","sliding_window"]}
def _cfg(strategy,chunk_size,overlap,preview_count=8):
    aliases={"table_aware":"table","code_aware":"code","sliding-window":"sliding_window","parent-child":"parent_child"}
    strategy=aliases.get(strategy,strategy)
    return ChunkingConfig(strategy=strategy,chunk_size=chunk_size,overlap=overlap,preview_count=preview_count)
@app.post("/v1/preview")
async def preview_document(request:Request,file:UploadFile=File(...),strategy:str=Form("auto"),chunk_size:int=Form(1000),overlap:int=Form(150),preview_count:int=Form(8)):
    if not file.filename:raise HTTPException(status_code=400,detail="Filename is required.")
    content=await file.read()
    try:
        result=engine.preview(file.filename,content,_cfg(strategy,chunk_size,overlap,preview_count),trace_id=request.state.trace_id).model_dump(mode="json")
        result["chunks"]=result.pop("preview_chunks"); result["total_chunks"]=result.pop("estimated_total_chunks"); result["is_preview"]=True
        return result
    except (ValueError,ValidationError) as exc:raise HTTPException(status_code=400,detail=str(exc)) from exc
    except Exception as exc:raise HTTPException(status_code=422,detail=f"Preview failed ({type(exc).__name__}). Check file format and logs using the trace ID.") from exc
    finally:await file.close()
@app.post("/v1/chunk")
async def chunk_document(request:Request,file:UploadFile=File(...),strategy:str=Form("auto"),chunk_size:int=Form(1000),overlap:int=Form(150)):
    if not file.filename:raise HTTPException(status_code=400,detail="Filename is required.")
    content=await file.read()
    try:return engine.process(file.filename,content,_cfg(strategy,chunk_size,overlap),trace_id=request.state.trace_id).model_dump(mode="json")
    except (ValueError,ValidationError) as exc:raise HTTPException(status_code=400,detail=str(exc)) from exc
    except Exception as exc:raise HTTPException(status_code=422,detail=f"Processing failed ({type(exc).__name__}). Check file format and logs using the trace ID.") from exc
    finally:await file.close()
class TextChunkRequest(BaseModel):
    text:str=Field(min_length=1); filename:str="inline.txt"; strategy:str="auto"; chunk_size:int=Field(default=1000,ge=100,le=10000); overlap:int=Field(default=150,ge=0,le=2000)
@app.post("/v1/chunk-text")
def chunk_text_endpoint(request:TextChunkRequest):
    try:return engine.process(request.filename,request.text.encode(),_cfg(request.strategy,request.chunk_size,request.overlap)).model_dump(mode="json")
    except (ValueError,ValidationError) as exc:raise HTTPException(status_code=400,detail=str(exc)) from exc


# Match Engine workflow: SDK chunking, OpenAI field extraction/embeddings, and Excel reconciliation.
import io, os, json, zipfile, math, base64, re, urllib.request, urllib.error
from difflib import SequenceMatcher
from openpyxl import load_workbook
import hashlib
import os, sqlite3, datetime

MATCH_FIELDS = ["case_number", "chain_code", "merchant_number", "transaction_amount", "avs_code", "reference_number"]
def _openai_post(api_key, endpoint, payload):
    req = urllib.request.Request("https://api.openai.com/v1/" + endpoint, data=json.dumps(payload).encode(), headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=60) as resp: return json.loads(resp.read().decode())
    except urllib.error.HTTPError as exc: raise HTTPException(status_code=502, detail="OpenAI API error: " + exc.read().decode(errors="replace")[:400]) from exc
    except Exception as exc: raise HTTPException(status_code=502, detail=f"OpenAI request failed: {exc}") from exc
def _norm(v): return re.sub(r"[^a-z0-9.]", "", str(v or "").strip().lower())
def _extract_fields(text, key, model):
    prompt = "Extract payment dispute fields from this text. Return JSON only with keys case_number, chain_code, merchant_number, transaction_amount, avs_code, reference_number. Use null if absent; preserve identifiers exactly; do not infer.\n\n" + text[:24000]
    data = _openai_post(key, "chat/completions", {"model": model, "temperature": 0, "response_format": {"type":"json_object"}, "messages":[{"role":"system","content":"Accurately extract structured values. JSON only."},{"role":"user","content":prompt}]})
    try: parsed = json.loads(data["choices"][0]["message"]["content"])
    except Exception as exc: raise HTTPException(status_code=502, detail="OpenAI returned invalid extraction JSON") from exc
    return {f: parsed.get(f) for f in MATCH_FIELDS}
def _embedding(text, key, model): return _openai_post(key, "embeddings", {"model":model,"input":text[:8000]})["data"][0]["embedding"]
def _cosine(a,b):
    dot=sum(x*y for x,y in zip(a,b)); na=math.sqrt(sum(x*x for x in a)); nb=math.sqrt(sum(y*y for y in b)); return dot/(na*nb) if na and nb else 0

def _case_row_text(row, cols):
    parts = []
    for field in MATCH_FIELDS:
        if field in cols:
            value = row["values"][cols[field]]
            if value not in (None, ""):
                parts.append(f"{field}: {value}")
    return " | ".join(parts)

def _chroma_collection(excel_filename, embedding_model):
    """Return a persistent, workbook/model-scoped Chroma collection."""
    try:
        import chromadb
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="ChromaDB is not installed. Install dependencies with: pip install -e '.[api]'") from exc
    db_path = os.getenv("CHROMA_DB_PATH", str(Path(__file__).resolve().parents[1] / "chroma_db"))
    Path(db_path).mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=db_path)
    identity = f"{Path(excel_filename).name.lower()}::{embedding_model}"
    suffix = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12]
    collection = client.get_or_create_collection(
        name=f"case_records_{suffix}",
        metadata={"hnsw:space": "cosine", "workbook": Path(excel_filename).name, "embedding_model": embedding_model},
    )
    return client, collection

def _sync_case_embeddings(collection, rows, cols, api_key, embedding_model):
    """Upsert changed case rows and remove records no longer present in Excel."""
    existing = collection.get(include=["metadatas", "documents"])
    existing_meta = {rid: ((existing.get("metadatas") or [])[i] or {}) for i, rid in enumerate(existing.get("ids") or [])}
    existing_ids = set(existing_meta)
    wanted_ids = set()
    for row in rows:
        rid = f"excel-row-{row['row']}"
        wanted_ids.add(rid)
        document = _case_row_text(row, cols)
        digest = hashlib.sha256(document.encode("utf-8")).hexdigest()
        old_meta = existing_meta.get(rid, {})
        if old_meta.get("content_hash") == digest and old_meta.get("embedding_model") == embedding_model:
            continue
        if not document.strip():
            continue
        vector = _embedding(document, api_key, embedding_model)
        metadata = {"excel_row": int(row["row"]), "case_number": str(row["case"] if row["case"] is not None else ""), "content_hash": digest, "embedding_model": embedding_model}
        collection.upsert(ids=[rid], embeddings=[vector], documents=[document], metadatas=[metadata])
    stale_ids = list(existing_ids - wanted_ids)
    if stale_ids:
        collection.delete(ids=stale_ids)

def _field_similarity(fields, row, cols):
    scores = []
    for field in MATCH_FIELDS:
        if field not in cols or fields.get(field) in (None, ""):
            continue
        value = row["values"][cols[field]]
        if value in (None, ""):
            continue
        left, right = _norm(fields[field]), _norm(value)
        scores.append(1.0 if left == right else 0.65 * SequenceMatcher(None, left, right).ratio())
    return sum(scores) / len(scores) if scores else 0.0


MATCH_DB_PATH = Path(os.getenv("MATCH_ENGINE_DB_PATH", str(BASE_DIR / "data" / "match_engine_history.sqlite3")))
CASE_WORKBOOK_PATH = Path(os.getenv("MATCH_ENGINE_CASE_WORKBOOK", str(BASE_DIR / "data" / "case_records.xlsx")))

def _history_conn():
    MATCH_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(MATCH_DB_PATH)
    conn.execute("""CREATE TABLE IF NOT EXISTS document_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT, document_name TEXT NOT NULL, file_type TEXT,
        pages INTEGER DEFAULT 0, status TEXT, strategy TEXT, uploaded_at TEXT,
        chunks INTEGER DEFAULT 0, quality_score REAL DEFAULT 0, chunk_size INTEGER DEFAULT 1000,
        error TEXT
    )""")
    conn.commit()
    return conn

def _record_history(name, file_type, pages, status, strategy, chunks, quality, chunk_size, error=None):
    conn = _history_conn()
    with conn:
        conn.execute("INSERT INTO document_history(document_name,file_type,pages,status,strategy,uploaded_at,chunks,quality_score,chunk_size,error) VALUES(?,?,?,?,?,?,?,?,?,?)",
            (name, file_type, pages, status, strategy, datetime.datetime.now(datetime.timezone.utc).isoformat(), chunks, quality, chunk_size, error))
    conn.close()

@app.get("/v1/match-engine/dashboard")
def match_engine_dashboard():
    conn = _history_conn()
    rows = conn.execute("SELECT document_name,file_type,pages,status,strategy,uploaded_at,chunks,quality_score,chunk_size,error FROM document_history ORDER BY id DESC LIMIT 100").fetchall()
    totals = conn.execute("SELECT COUNT(*), COALESCE(SUM(chunks),0), COALESCE(AVG(quality_score),0) FROM document_history").fetchone()
    trend_rows = conn.execute("SELECT substr(uploaded_at,1,10), ROUND(AVG(quality_score),1), COUNT(*) FROM document_history GROUP BY substr(uploaded_at,1,10) ORDER BY substr(uploaded_at,1,10) DESC LIMIT 7").fetchall()
    distribution = conn.execute("SELECT strategy, COUNT(*) FROM document_history GROUP BY strategy").fetchall()
    conn.close()
    recent = [{"document_name":r[0],"type":r[1],"pages":r[2],"status":r[3],"strategy":r[4],"uploaded_at":r[5],"chunks":r[6],"quality_score":r[7],"chunk_size":r[8],"error":r[9]} for r in rows]
    case_count = 0
    if CASE_WORKBOOK_PATH.exists():
        try:
            case_wb = load_workbook(CASE_WORKBOOK_PATH, read_only=True, data_only=True)
            case_ws = case_wb.active
            case_count = max(0, case_ws.max_row - 1)
            case_wb.close()
        except Exception:
            case_count = 0
    return {"total_documents":totals[0],"total_chunks":totals[1],"average_quality":round(totals[2],1),"recent_documents":recent,
        "quality_trend":[{"date":r[0],"quality":r[1],"documents":r[2]} for r in reversed(trend_rows)],
        "chunk_distribution":[{"strategy":r[0] or "unknown","count":r[1]} for r in distribution],
        "processing_pipeline":[{"name":"Upload","status":"ready"},{"name":"SDK chunking","status":"ready"},{"name":"Field extraction","status":"ready"},{"name":"OpenAI embeddings","status":"backend configured"},{"name":"ChromaDB matching","status":"persistent"}],
        "case_database":{"filename":CASE_WORKBOOK_PATH.name,"available":CASE_WORKBOOK_PATH.exists(),"records":case_count}}

@app.post("/v1/match-engine")
async def match_engine(request: Request, documents: list[UploadFile] = File(...), strategy: str = Form("auto"), chunk_size: int = Form(1000), overlap: int = Form(150)):
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    llm_model = os.getenv("MATCH_ENGINE_LLM_MODEL", "gpt-5")
    embedding_model = os.getenv("MATCH_ENGINE_EMBEDDING_MODEL", "text-embedding-3-small")
    if not api_key: raise HTTPException(status_code=500, detail="OpenAI is not configured on the backend. Set OPENAI_API_KEY.")
    if not CASE_WORKBOOK_PATH.exists(): raise HTTPException(status_code=500, detail=f"Case database workbook not found: {CASE_WORKBOOK_PATH}")
    try: wb=load_workbook(CASE_WORKBOOK_PATH); ws=wb.active
    except Exception as exc: raise HTTPException(status_code=500, detail=f"Cannot read project case database: {exc}") from exc
    excel_filename = CASE_WORKBOOK_PATH.name
    headers=[str(c.value or "").strip() for c in ws[1]]; norm_headers=[_norm(h) for h in headers]
    aliases={"case_number":["casenumber","case","caseno"],"chain_code":["chaincode","chain","chainid"],"merchant_number":["merchantnumber","merchantno","merchant","mid"],"transaction_amount":["transactionamount","amount","txnamount"],"avs_code":["avscode","avs"],"reference_number":["referencenumber","reference","arn","retrievalreferencenumber"],"document_name":["documentname","document","filename","sourcefilename"]}
    cols={}
    for field, names in aliases.items():
        for i,h in enumerate(norm_headers):
            if h in names: cols[field]=i; break
    if "case_number" not in cols: raise HTTPException(status_code=400, detail="Excel must have a Case Number column")
    rows=[]
    for rn in range(2,ws.max_row+1):
        vals=[ws.cell(rn,c+1).value for c in range(len(headers))]
        if any(v is not None for v in vals): rows.append({"row":rn,"values":vals,"case":vals[cols["case_number"]]})
    if not rows: raise HTTPException(status_code=400, detail="No case rows found in Excel")
    # Persist case vectors in ChromaDB, scoped to this workbook and embedding model.
    client, collection = _chroma_collection(excel_filename, embedding_model)
    _sync_case_embeddings(collection, rows, cols, api_key, embedding_model)
    rows_by_number = {row["row"]: row for row in rows}
    results=[]
    for upload in documents:
        fname=upload.filename or "document.pdf"; raw=await upload.read(); pdfs=[]
        if fname.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                    pdfs=[(n,z.read(n)) for n in z.namelist() if not n.endswith("/") and n.lower().endswith(".pdf")]
            except Exception as exc: results.append({"document":fname,"status":"error","error":f"Invalid ZIP: {exc}"}); continue
        elif fname.lower().endswith(".pdf"): pdfs=[(fname,raw)]
        else: results.append({"document":fname,"status":"skipped","error":"Upload a PDF or ZIP containing PDFs"}); continue
        for docname, content in pdfs:
            try:
                chunked=engine.process(docname,content,_cfg(strategy,chunk_size,overlap),trace_id=request.state.trace_id)
                try:
                    from pypdf import PdfReader
                    page_count = len(PdfReader(io.BytesIO(content)).pages)
                except Exception:
                    page_count = 0
                chunk_lengths = [len(c.text.strip()) for c in chunked.chunks if c.text.strip()]
                mean_len = (sum(chunk_lengths) / len(chunk_lengths)) if chunk_lengths else 0
                # Heuristic dashboard quality score: rewards non-empty, reasonably balanced chunks.
                balance_penalty = min(1.0, max(0.0, (max(chunk_lengths, default=0) - mean_len) / max(1.0, mean_len * 4)))
                nonempty_ratio = len(chunk_lengths) / max(1, len(chunked.chunks))
                quality_score = round(max(0.0, min(100.0, 100.0 * nonempty_ratio * (1.0 - balance_penalty))), 1)
                text="\n\n".join(c.text for c in chunked.chunks)
                fields=_extract_fields(text,api_key,llm_model)
                doc_emb=_embedding(" ".join(f"{k}: {v}" for k,v in fields.items() if v not in (None,""))+"\n"+text[:4000],api_key,embedding_model)
                found=None; method="unmatched"; confidence=0.0
                case=_norm(fields.get("case_number"))
                if case:
                    found=next((r for r in rows if _norm(r["case"])==case),None)
                    if found: method="exact_case_number"; confidence=1.0
                if not found:
                    # ChromaDB performs persistent vector search. Strong field matches are also
                    # considered so a useful exact identifier match is not lost outside top-k.
                    count = collection.count()
                    vector_hits = collection.query(query_embeddings=[doc_emb], n_results=min(25, count), include=["metadatas", "distances"]) if count else {"ids": [[]], "distances": [[]], "metadatas": [[]]}
                    semantic_scores = {}
                    for rid, distance in zip(vector_hits.get("ids", [[]])[0], vector_hits.get("distances", [[]])[0]):
                        try:
                            row_num = int(rid.removeprefix("excel-row-"))
                            semantic_scores[row_num] = max(0.0, min(1.0, 1.0 - float(distance)))
                        except (ValueError, TypeError):
                            continue
                    candidate_rows = {n: rows_by_number[n] for n in semantic_scores if n in rows_by_number}
                    for row in rows:
                        if _field_similarity(fields, row, cols) >= .35:
                            candidate_rows[row["row"]] = row
                    ranked=[]
                    for row in candidate_rows.values():
                        field_score = _field_similarity(fields, row, cols)
                        semantic = semantic_scores.get(row["row"], 0.0)
                        if row["row"] not in semantic_scores:
                            row_text = _case_row_text(row, cols)
                            row_vector = _embedding(row_text, api_key, embedding_model) if row_text.strip() else []
                            semantic = _cosine(doc_emb, row_vector)
                        combined = .8 * field_score + .2 * semantic
                        if field_score >= .35 or semantic >= .45:
                            ranked.append((combined, row))
                    if ranked:
                        confidence, candidate=max(ranked,key=lambda x:x[0])
                        if confidence >= .55: found=candidate; method="chroma_vector_and_field_similarity"
                        else: confidence=0
                if found:
                    if "document_name" not in cols:
                        cols["document_name"]=len(headers); ws.cell(1,len(headers)+1).value="Document Name"; headers.append("Document Name")
                    ws.cell(found["row"],cols["document_name"]+1).value=docname
                results.append({"document":docname,"status":"matched" if found else "unmatched","extracted_fields":fields,"matched_case_number":found["case"] if found else None,"excel_row":found["row"] if found else None,"match_method":method,"confidence":round(confidence,4),"chunks_created":len(chunked.chunks),"sdk_strategy":chunked.config.strategy,"pages":page_count,"quality_score":quality_score})
            except HTTPException: raise
            except Exception as exc: results.append({"document":docname,"status":"error","error":f"{type(exc).__name__}: {exc}"})
    wb.save(CASE_WORKBOOK_PATH)
    for item in results:
        status=item.get("status","error")
        try:
            _record_history(item.get("document","unknown"), Path(item.get("document","")).suffix.lower().lstrip(".") or "pdf",
                int(item.get("pages",0) or 0), status, item.get("sdk_strategy",strategy), int(item.get("chunks_created",0) or 0),
                float(item.get("quality_score",0) or 0), chunk_size, item.get("error"))
        except Exception:
            pass
    return {"results":results,"summary":{"processed":len(results),"matched":sum(r.get("status")=="matched" for r in results),"unmatched":sum(r.get("status")=="unmatched" for r in results),"errors":sum(r.get("status")=="error" for r in results)},"case_database":CASE_WORKBOOK_PATH.name,"embedding_store":{"type":"ChromaDB","collection":collection.name,"persist_directory":os.getenv("CHROMA_DB_PATH",str(Path(__file__).resolve().parents[1]/"chroma_db")),"records":collection.count(),"persistent":True}}
